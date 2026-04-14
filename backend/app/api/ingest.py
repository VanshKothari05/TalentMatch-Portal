from __future__ import annotations

import csv
import io
import json
import logging
import re
from typing import List

from fastapi import APIRouter, HTTPException, UploadFile, File

from app.core.embeddings import embedding_service
from app.core.matcher import CandidateProfile, JobDescription
from app.core.store import store
from app.models.schemas import CandidateIn, IngestResponse, JobIn

logger = logging.getLogger(__name__)
router = APIRouter(tags=["Ingest"])

# ── Helpers ───────────────────────────────────────────────────────────────────

KNOWN_SKILLS = [
    "python","java","javascript","typescript","go","rust","ruby","php","c++","c#","scala",
    "django","fastapi","flask","spring boot","express.js","rails","laravel","next.js",
    "react","vue","angular","tailwind css","html","css",
    "postgresql","mysql","mongodb","redis","elasticsearch","cassandra","sqlite",
    "docker","kubernetes","aws","gcp","azure","terraform","ci/cd","jenkins","github actions",
    "tensorflow","pytorch","scikit-learn","mlflow","machine learning","deep learning",
    "kafka","rabbitmq","celery","sqs","rest apis","graphql","grpc","microservices",
    "sql","nosql","git","linux","bash","airflow","llm","rag","embeddings","langchain",
    "openai","vector database","pinecone","weaviate","chroma","fastapi","pydantic",
    "prompt engineering","fine-tuning","hugging face","bert","gpt",
]

def _extract_skills(text: str) -> List[str]:
    text_lower = text.lower()
    return [s for s in KNOWN_SKILLS if re.search(r"\b" + re.escape(s) + r"\b", text_lower)]

def _safe_int(v, default=0):
    try:
        return int(float(str(v).strip()))
    except:
        return default

def _parse_skills(raw: str) -> List[str]:
    if not raw:
        return []
    return [s.strip() for s in raw.split(",") if s.strip()]

def _read_pdf_bytes(data: bytes) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [p.extract_text() or "" for p in reader.pages]
        raw = "\n".join(pages)
        return _clean_text(raw)
    except Exception as e:
        logger.error("PDF read error: %s", e)
        return ""

def _read_docx_bytes(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        return _clean_text("\n".join(p.text for p in doc.paragraphs))
    except Exception as e:
        logger.error("DOCX read error: %s", e)
        return ""

def _clean_text(raw: str) -> str:
    import re
    text = raw.replace("\n \n", " ")
    text = re.sub(r"\n +\n", " ", text)
    text = re.sub(r" {2,}", " ", text)
    text = re.sub(r"\. ([A-Z])", ".\n\n\\1", text)
    BULLETS = ["\u25cf","\uf0b7","\uf0a7","\u25a1","\u25aa","\u25ab",
               "\u2022","\u2023","\u2043","\u204c","●","•","·","▪","▸","►"]
    for b in BULLETS:
        text = text.replace(b, "\x00BULLET\x00")
    parts = text.split("\x00BULLET\x00")
    result = []
    HEADING_ENDINGS = {"requirements","qualifications","responsibilities","overview",
                       "benefits","compensation","skills","guidelines","expectations","environment"}
    for i, part in enumerate(parts):
        part = part.strip()
        if not part:
            continue
        if i == 0:
            result.append(part)
            continue
        first_line = part.split(".")[0].split("\n")[0].strip()
        words = first_line.split()
        is_heading = (
            1 <= len(words) <= 4
            and all(w[0].isupper() for w in words if w.isalpha())
            and words[-1].lower().rstrip("s") in {h.rstrip("s") for h in HEADING_ENDINGS}
            and not first_line.endswith(",")
        )
        result.append(("\n\n" if is_heading else "\n• ") + part)
    text = "".join(result)
    text = re.sub(r" {2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _parse_csv_candidates(data: bytes) -> List[CandidateProfile]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    profiles = []
    for row in reader:
        skills_raw = (row.get("parsed_skills") or row.get("programming_languages")
                      or row.get("skills") or "")
        cid = str(row.get("id", "")).strip()
        if not cid:
            continue
        profiles.append(CandidateProfile(
            id=cid,
            name=row.get("name", "").strip(),
            skills=_parse_skills(skills_raw),
            years_experience=_safe_int(row.get("years_of_experience", 0)),
            summary=row.get("parsed_summary", "") or "",
            education=row.get("degree", "").strip(),
            current_role=row.get("current_title", "").strip(),
        ))
    return profiles

def _parse_excel_candidates(data: bytes) -> List[CandidateProfile]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip().lower() if h else "" for h in rows[0]]
        profiles = []
        for row in rows[1:]:
            d = {headers[i]: (str(row[i]).strip() if row[i] is not None else "") for i in range(min(len(headers), len(row)))}
            skills_raw = d.get("parsed_skills") or d.get("skills") or d.get("programming_languages") or ""
            cid = d.get("id", "").strip()
            if not cid:
                continue
            profiles.append(CandidateProfile(
                id=cid,
                name=d.get("name", ""),
                skills=_parse_skills(skills_raw),
                years_experience=_safe_int(d.get("years_of_experience", 0)),
                summary=d.get("parsed_summary", "") or d.get("summary", ""),
                education=d.get("degree", "") or d.get("education", ""),
                current_role=d.get("current_title", "") or d.get("current_role", ""),
            ))
        return profiles
    except Exception as e:
        logger.error("Excel read error: %s", e)
        return []

# ── Job upload ────────────────────────────────────────────────────────────────

@router.post("/jobs", response_model=IngestResponse)
async def upload_jobs(jobs: List[JobIn]):
    ids = []
    for j in jobs:
        store.add_job(JobDescription(
            id=j.id, title=j.title, company=j.company,
            description=j.description, required_skills=j.required_skills,
            preferred_skills=j.preferred_skills, min_experience=j.min_experience,
            max_experience=j.max_experience, hiring_mode=j.hiring_mode,
        ))
        ids.append(j.id)
    return IngestResponse(status="ok", added=len(ids), ids=ids)


@router.post("/jobs/upload", response_model=IngestResponse)
async def upload_job_file(file: UploadFile = File(...)):
    data = await file.read()
    fname = (file.filename or "").lower()

    if fname.endswith(".pdf"):
        text = _read_pdf_bytes(data)
    elif fname.endswith(".docx"):
        text = _read_docx_bytes(data)
    elif fname.endswith(".txt"):
        text = data.decode("utf-8", errors="replace").strip()
    elif fname.endswith(".json"):
        parsed = json.loads(data)
        if isinstance(parsed, list):
            return await upload_jobs([JobIn(**i) for i in parsed])
        text = parsed.get("description", str(parsed))
    else:
        text = data.decode("utf-8", errors="replace").strip()

    if not text:
        raise HTTPException(status_code=400, detail="Could not extract text from file.")

    skills = _extract_skills(text)
    split = max(1, len(skills) // 2)

    import hashlib
    jd_id = "jd_" + hashlib.md5(text[:200].encode()).hexdigest()[:8]

    jd = JobDescription(
        id=jd_id, title="Uploaded Job Role", company="",
        description=text, required_skills=skills[:split],
        preferred_skills=skills[split:], min_experience=0,
        max_experience=20, hiring_mode="default",
    )
    store.add_job(jd)
    return IngestResponse(status="ok", added=1, ids=[jd_id])


# ── Candidate upload ──────────────────────────────────────────────────────────

@router.post("/candidates", response_model=IngestResponse)
async def upload_candidates(candidates: List[CandidateIn]):
    ids = []
    texts = {}
    for c in candidates:
        profile = CandidateProfile(
            id=c.id, name=c.name, skills=c.skills,
            years_experience=c.years_experience, summary=c.summary,
            education=c.education, current_role=c.current_role,
        )
        store.add_candidate(profile)
        texts[c.id] = profile.full_text
        ids.append(c.id)
    embedding_service.add_candidates(texts)
    return IngestResponse(status="ok", added=len(ids), ids=ids)


@router.post("/candidates/upload", response_model=IngestResponse)
async def upload_candidates_file(file: UploadFile = File(...)):
    data = await file.read()
    fname = (file.filename or "").lower()
    profiles: List[CandidateProfile] = []

    if fname.endswith(".csv"):
        profiles = _parse_csv_candidates(data)
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        profiles = _parse_excel_candidates(data)
    elif fname.endswith(".json"):
        parsed = json.loads(data)
        return await upload_candidates([CandidateIn(**i) for i in parsed])
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {fname}. Use CSV, Excel, or JSON.")

    if not profiles:
        raise HTTPException(status_code=400, detail="No valid candidates found in file.")

    texts = {}
    for p in profiles:
        store.add_candidate(p)
        texts[p.id] = p.full_text

    embedding_service.add_candidates(texts)
    return IngestResponse(status="ok", added=len(profiles), ids=[p.id for p in profiles])